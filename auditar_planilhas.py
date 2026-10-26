"""Auditoria de leitura: preserva os XLSX e gera CSV/JSON rastreáveis."""
import argparse
from collections import Counter, defaultdict
import csv
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from data_logic import (CAT_FIELDS, INI_FIELDS, CATEGORIES, clean, header,
                        load_collection, metadata_count, normalize)


def audit(folder, output):
    folder, output = Path(folder), Path(output)
    output.mkdir(parents=True, exist_ok=True)
    paths = sorted(folder.glob("BR-*.xlsx"))
    cat, ini = load_collection([p.name for p in paths], folder)
    findings, summary, snapshots = [], [], []

    def note(filename, sheet, cell, problem, original, action):
        findings.append({"Arquivo": filename, "Aba": sheet, "Célula": cell,
                         "Problema": problem, "Valor original": str(original or ""),
                         "Encaminhamento": action})

    for path in paths:
        book = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        trims = dummy = formula_errors = blanks = 0
        source_records = source_metadata = reps = 0
        for sheet in book:
            rows = list(sheet.iter_rows())
            cached_rows = list(cached[sheet.title].iter_rows())
            values = [[c.value for c in row] for row in cached_rows]
            snapshots.append({"arquivo": path.name, "aba": sheet.title, "linhas": values})
            for row_idx, row in enumerate(rows):
                for col_idx, cell in enumerate(row):
                    value = cell.value
                    if value is None:
                        continue
                    cv = cached_rows[row_idx][col_idx] if row_idx < len(cached_rows) and col_idx < len(cached_rows[row_idx]) else None
                    if cell.data_type == "f":
                        if "__xludf.DUMMYFUNCTION" in value:
                            dummy += 1
                            note(path.name, sheet.title, cell.coordinate, "Fórmula exportada do Google Sheets", value,
                                 "O XLSX contém resultado salvo. Recalcular no Google Sheets antes de exportar; não substituir o resultado por zero.")
                        if cv is not None and cv.data_type == "e":
                            formula_errors += 1
                            note(path.name, sheet.title, cell.coordinate, "Erro no resultado salvo da fórmula", cv.value,
                                 "Revisar a fórmula de origem. O erro não é um metadado preenchido.")
                        elif cv is None or cv.value is None:
                            # Muitas fórmulas em linhas-modelo retornam vazio legitimamente.
                            if sheet.title == "Geral" and row_idx >= 2 and len(cached_rows[row_idx]) > 1 and clean(cached_rows[row_idx][1].value):
                                note(path.name, sheet.title, cell.coordinate, "Fórmula sem resultado salvo em registro preenchido", value,
                                     "Conferir se o campo é opcional ou se precisa recalcular na origem.")
                    elif isinstance(value, str):
                        prepared = re.sub(r"[\u200b-\u200d\u2060\ufeff]", "", value.replace("\xa0", " ")).strip()
                        if value != prepared:
                            trims += 1
                            note(path.name, sheet.title, cell.coordinate, "Espaço externo ou caractere invisível", value,
                                 "Removido apenas na cópia tabular de análise; original preservado.")

            if sheet.title == "Geral" and "MAPEAMENTOS" in path.name:
                for idx, row in enumerate(values[2:], 3):
                    if len(row) > 1 and clean(row[1]):
                        source_records += 1
                        source_metadata += sum(bool(clean(v)) for v in row[1:20 if "MSSCPENHA" in path.name else 19])
                        title_col = 10 if "MSSCPENHA" in path.name else 9
                        if not clean(row[title_col]):
                            note(path.name, sheet.title, f"{get_column_letter(title_col+1)}{idx}", "Registro sem título do documento", row[title_col],
                                 "Preservado por possuir iniciativa preenchida; completar o título na origem.")
                    elif any(clean(v) for v in row):
                        blanks += 1
                        note(path.name, sheet.title, f"B{idx}", "Linha sem iniciativa", row[1] if len(row)>1 else "",
                             "Excluída da cópia de análise. Preservada no original como linha incompleta ou modelo.")
            elif sheet.title == "Geral":
                reps += sum(any(clean(row[j]) for j in (10,12,14,16) if j < len(row)) for row in values[1:])
            elif values and len(values[0]) > 1 and "titulo descritivo" in normalize(values[0][1]):
                data = [(idx,row) for idx,row in enumerate(values[1:],2) if len(row)>1 and clean(row[1])]
                source_records += len(data)
                source_metadata += sum(bool(clean(v)) for _,row in data for v in row[:20])
                for idx, row in data:
                    technique = clean(row[16]) if len(row)>16 else ""
                    code = clean(row[10]) if len(row)>10 else ""
                    classification = re.search(r"_([A-Z0-9]+)-([A-Z0-9]+)-([A-Z0-9]+)-([A-Z0-9]+)_", code)
                    if "digitalizacao" in normalize(technique) and classification and classification.group(3) == "FOT":
                        note(path.name, sheet.title, f"Q{idx}", "Técnica detalhada diverge do código", technique,
                             f"Descrição informa digitalização, código contém FOT: {code}. Conferir também Geral!O. Não reclassificado automaticamente.")
        kind = "initiatives" if "MAPEAMENTOS" in path.name else "catalogue"
        frame = ini if kind == "initiatives" else cat
        frame = frame[frame.Arquivo_origem.eq(path.name)]
        assert len(frame) == source_records, (path.name, len(frame), source_records)
        assert metadata_count(frame, kind) == source_metadata, (path.name, source_metadata)
        summary.append({"Arquivo":path.name, "Registros": source_records,
                        "Representações": reps if kind == "catalogue" else None,
                        "Metadados de origem":source_metadata, "Linhas sem iniciativa":blanks,
                        "Espaços e invisíveis":trims, "Fórmulas Google Sheets":dummy,
                        "Erros de fórmula salvos":formula_errors,
                        "SHA256":hashlib.sha256(path.read_bytes()).hexdigest()})
        book.close(); cached.close()

    for filename, group in cat.groupby("Arquivo_origem"):
        for code, repeated in group.groupby("Código de referência"):
            if code and len(repeated)>1:
                for _,record in repeated.iterrows():
                    note(filename, record.Aba_origem, f"K{record.Linha_origem}", "Código de referência repetido entre descrições", code,
                         "Preservado. Validar cada código com a Geral. As descrições são distintas e não devem ser apagadas como duplicatas.")

    def tabular(frame):
        rows=[]
        for record in frame.to_dict("records"):
            rows.append({k:("; ".join(v) if isinstance(v,(tuple,list)) else v) for k,v in record.items() if k not in {"Representações_ID","Conflito de código","Representações vinculadas"}})
        return rows

    catalogue = tabular(cat)
    mappings = tabular(ini)
    representations = [{k:v for k,v in r.items() if k not in {"Chave", "Chave documental"}} for r in cat.attrs["representations"]]
    for name,records in [("inventario",catalogue),("mapeamentos",mappings),("representacoes",representations),("checagens",findings),("resumo",summary)]:
        with (output/f"{name}.csv").open("w",encoding="utf-8-sig",newline="") as f:
            writer=csv.DictWriter(f,fieldnames=list(records[0]) if records else [])
            writer.writeheader(); writer.writerows(records)
    payload={"resumo":summary,"inventario":catalogue,"mapeamentos":mappings,"representacoes":representations,"checagens":findings}
    (output/"auditoria.json").write_text(json.dumps(payload,ensure_ascii=False,default=str),encoding="utf-8")
    print(json.dumps({"registros_inventario":len(cat),"registros_mapeados":len(ini),"representacoes":len(representations),
                      "metadados":sum(s["Metadados de origem"] for s in summary),
                      "fórmulas_google":sum(s["Fórmulas Google Sheets"] for s in summary),
                      "erros_salvos":sum(s["Erros de fórmula salvos"] for s in summary),
                      "espacos_invisiveis":sum(s["Espaços e invisíveis"] for s in summary),
                      "problemas":dict(Counter(f["Problema"] for f in findings))},ensure_ascii=False))
    return payload


if __name__ == "__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pasta",default=str(Path(__file__).resolve().parent))
    parser.add_argument("--saida",default="auditoria-planilhas")
    args=parser.parse_args()
    audit(args.pasta,args.saida)
