"""Interface das duas áreas de análise; os cálculos ficam em data_logic."""
from io import BytesIO
import re

import matplotlib.pyplot as plt
import pandas as pd
import plotly.express as px
import streamlit as st
from wordcloud import WordCloud
from docx import Document

from data_logic import (CATEGORIES, CAT_FIELDS, INI_FIELDS, category_options,
                        display_table, filter_categories, metadata_count,
                        normalize, phrase_frequency, search_records, year_frequency)
from vocabulario_controlado import rotular_sigla

ALL_CARANDIRU = "Rememorações todas (Carandiru)"
PENHA_SUFFIX = " (Massacre da Penha)"


def _reset(prefix, signature):
    if st.session_state.get(prefix + "signature") != signature:
        for key in list(st.session_state):
            if key.startswith(prefix) and key != prefix + "files":
                del st.session_state[key]
        st.session_state[prefix + "signature"] = signature


def _intervention_changed(prefix):
    key = prefix + "Intervenção"
    selected = st.session_state.get(key, [])
    previous = st.session_state.get(key + "_previous", [])
    added = set(selected) - set(previous)
    if ALL_CARANDIRU in added:
        selected = [v for v in selected if v == ALL_CARANDIRU or v.endswith(PENHA_SUFFIX)]
    elif any(not v.endswith(PENHA_SUFFIX) and v != ALL_CARANDIRU for v in added):
        selected = [v for v in selected if v != ALL_CARANDIRU]
    st.session_state[key] = selected
    st.session_state[key + "_previous"] = selected.copy()


def initiative_filter(frame, selected):
    if not selected:
        return frame
    car = frame["Arquivo_origem"].str.contains("REMEMORA-CARANDIRU", regex=False)
    pen = frame["Arquivo_origem"].str.contains("MSSCPENHA", regex=False)
    mask = pd.Series(False, index=frame.index)
    for value in selected:
        if value == ALL_CARANDIRU:
            mask |= car
        elif value.endswith(PENHA_SUFFIX):
            mask |= pen & (frame["Finalidade primária"] == value[:-len(PENHA_SUFFIX)])
        else:
            mask |= car & (frame["Intervenção"] == value)
    return frame.loc[mask].copy()


def _export_docx(frame):
    doc = Document()
    doc.add_heading("Inventário filtrado", 0)
    doc.add_paragraph(f"Registros: {len(frame)}")
    for _, row in frame.iterrows():
        title = row.get("Título (Busca)", row.get("Título do documento", "")) or row.get("Nome da iniciativa", "Registro")
        doc.add_heading(str(title), 2)
        for field, value in row.items():
            if str(value).strip() and field not in {"Registro_ID", "Representações_ID", "Conflito de código"}:
                doc.add_paragraph(f"{field}: {value}")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_analysis(kind, available, descriptions, loader, folder, translate):
    prefix = "cat_" if kind == "catalogue" else "inic_"
    files = st.multiselect(translate("Selecione as planilhas para integrar:"), available, default=available, key=prefix+"files")
    _reset(prefix, tuple(sorted(files)))
    for filename in files:
        description = descriptions.get(filename, {}).get(st.session_state.get("app_language", "Português"), "")
        if description:
            st.caption(f"{filename}: {description}")
    if not files:
        st.info("Selecione ao menos uma planilha nesta aba.")
        return
    cat, ini = loader(files, folder)
    original = cat if kind == "catalogue" else ini
    frame = original
    if kind == "catalogue":
        unit = st.radio("Unidade da contagem", ["Descrições documentais", "Representações / arquivos (Geral)"], horizontal=True, key=prefix+"unit")
        if unit.startswith("Representações"):
            frame = pd.DataFrame(cat.attrs.get("representations", []))
            if frame.empty:
                st.info("Não há representações nas bases selecionadas.")
                return
        else:
            st.caption("Uma descrição pode ter várias representações, como frente e verso. Códigos conflitantes são preservados e sinalizados.")
    st.subheader(translate("Busca avançada"))
    term = st.text_input("Pesquisar palavra, número ou frase inteira", key=prefix+"search")
    if "Referência do arquivo" in frame:
        filtered = frame[frame["Referência do arquivo"].map(normalize).str.contains(normalize(term), regex=False)] if term else frame.copy()
    else:
        filtered = search_records(frame, term)
    st.subheader(translate("Filtros categoriais"))
    fields = list(CATEGORIES) if kind == "catalogue" else ["Nome da iniciativa", "Intervenção", "Abrangência", "Modalidade"]
    columns = st.columns(4)
    selections = {}
    for column, field in zip(columns, fields):
        with column:
            key = prefix + field
            if field == "Intervenção":
                car = ini[ini["Arquivo_origem"].str.contains("REMEMORA-CARANDIRU", regex=False)]
                pen = ini[ini["Arquivo_origem"].str.contains("MSSCPENHA", regex=False)]
                options = ([ALL_CARANDIRU] if not car.empty else []) + category_options(car, field) + [v+PENHA_SUFFIX for v in category_options(pen, "Finalidade primária")]
                default = [ALL_CARANDIRU] if not car.empty and pen.empty else []
                if key not in st.session_state:
                    st.session_state[key] = default
                selected = st.pills(translate(field), options, selection_mode="multi", key=key, on_change=_intervention_changed, args=(prefix,), help="Carandiru: intervenção. Penha: finalidade primária. Os dois campos originais são preservados na tabela.")
                filtered = initiative_filter(filtered, selected)
            else:
                options = category_options(frame, field)
                if key in st.session_state:
                    st.session_state[key] = [v for v in st.session_state[key] if v in options]
                def label(value, field=field):
                    if kind == "catalogue":
                        type_key = dict(zip(CATEGORIES, ("genero", "especie", "tecnica", "forma")))[field]
                        return rotular_sigla(value, tipo=type_key)
                    sources = set(ini.loc[ini[field] == value, "Arquivo_origem"])
                    return value + PENHA_SUFFIX if sources and all("MSSCPENHA" in f for f in sources) else value
                selections[field] = st.multiselect(translate(field), options, key=key, format_func=label, help='Este campo é apenas o detalhamento do campo "Abrangência".' if field == "Modalidade" else None)
    filtered = filter_categories(filtered, selections)
    st.subheader(translate("Indicadores"))
    metrics = st.columns(3)
    metrics[0].metric("Registros exibidos", len(filtered))
    if kind == "catalogue":
        metrics[1].metric("Gêneros documentais", len(category_options(filtered, "Gênero documental")))
    else:
        metrics[1].metric("Bases selecionadas", len(files))
    is_representation = "Representação_ID" in filtered
    count = sum(bool(str(v).strip()) for f in CATEGORIES for v in filtered[f]) if is_representation else metadata_count(filtered, kind)
    metrics[2].metric("Metadados preenchidos no recorte", count)
    if is_representation:
        st.caption("Na visão de representações, metadados preenchidos contam as quatro classificações da Geral.")
    else:
        st.caption("Metadados preenchidos contam os campos de origem, uma vez por registro. Identificadores internos e classificações derivadas da Geral não entram novamente na soma.")
    st.subheader(translate("Análises e visualizações do acervo"))
    views = ["Nenhuma visualização (limpar tela)", "Linha do tempo (distribuição cronológica)", "Frequências categoriais", "Nuvem de palavras"]
    if kind == "catalogue" and not is_representation:
        views.append("Frequências temáticas")
    visualization = st.selectbox("Escolha uma visualização", views, key=prefix+"view")
    if filtered.empty:
        st.info("Nenhum registro corresponde ao recorte selecionado.")
    elif visualization == views[1]:
        field = "Data (Busca)" if kind == "catalogue" else "Ano"
        if field not in filtered:
            st.info("A Geral não informa datas. Selecione Descrições documentais para analisar a cronologia.")
        else:
            frequency, dated, undated = year_frequency(filtered, field)
            st.caption(f"Registros com ano: {dated}. Sem ano extraível: {undated}. Anos sem registros no intervalo aparecem com zero.")
            if kind == "catalogue":
                st.caption("As fontes utilizam datas de criação ou difusão. O gráfico apresenta a data registrada em cada descrição.")
            if not frequency.empty:
                fig = px.line(frequency, x="Ano", y="Frequência", markers=True)
                fig.update_layout(yaxis_title="Registros", xaxis_title="Ano")
                st.plotly_chart(fig, width="stretch")
                st.dataframe(frequency, hide_index=True, width="stretch")
    elif visualization == views[2]:
        field = st.selectbox("Campo da frequência", fields, key=prefix+"frequency_field")
        counts = {v: int(filter_categories(filtered, {field:[v]}).shape[0]) for v in category_options(filtered, field)}
        frequency = pd.DataFrame(list(counts.items()), columns=["Categoria", "Registros"])
        st.caption("Cada registro conta uma vez por categoria. Um registro com várias categorias pode entrar em mais de uma barra.")
        st.dataframe(frequency, hide_index=True, width="stretch")
        if not frequency.empty:
            st.plotly_chart(px.bar(frequency, x="Categoria", y="Registros"), width="stretch")
    elif visualization == views[3]:
        options = ["Título (Busca)", "Conteúdo (Busca)", "Palavras-chave"] if kind == "catalogue" else ["Nome da iniciativa", "Ano", "Fonte / Origem", "Proponente", "Título do documento"]
        options = [f for f in options if f in filtered]
        chosen = st.multiselect("O que deve conter?", options, default=options, key=prefix+"cloud_fields")
        combined, detail = phrase_frequency(filtered, chosen)
        st.caption("Frequência somada nos campos selecionados. Vírgulas separam termos; espaços preservam nomes e expressões completas.")
        if combined:
            try:
                cloud = WordCloud(width=1000, height=450, background_color="rgba(0,0,0,0)", mode="RGBA", collocations=False, max_words=150, random_state=7).generate_from_frequencies(combined)
                fig, ax = plt.subplots(figsize=(12, 5))
                ax.imshow(cloud, interpolation="bilinear")
                ax.axis("off")
                st.pyplot(fig)
                plt.close(fig)
            except ValueError:
                st.info("Os termos não couberam na nuvem. As frequências completas continuam disponíveis na tabela.")
            st.dataframe(detail.sort_values("Frequência", ascending=False), hide_index=True, width="stretch")
        else:
            st.info("Selecione campos com termos preenchidos para gerar a nuvem.")
    elif visualization == "Frequências temáticas":
        themes = {"Família":["mãe", "filho", "criança", "pai", "avó"], "Educação, artes e ofícios":["escola", "alfabetização", "atividade cultural", "costura"], "Arquitetura prisional":["grade", "cela", "pavilhão", "parede", "portão"]}
        theme = st.selectbox("Tema", list(themes), key=prefix+"theme")
        frequency = pd.DataFrame([{"Termo": term, "Registros": len(search_records(filtered[["Título (Busca)", "Conteúdo (Busca)"]], term))} for term in themes[theme]])
        st.caption("Documentos contendo cada termo, independentemente do número de repetições. Um documento pode conter vários termos.")
        st.plotly_chart(px.bar(frequency, x="Termo", y="Registros"), width="stretch")
        st.dataframe(frequency, hide_index=True)
    if visualization != views[0]:
        with st.expander("Conferir registros e exportar"):
            view = display_table(filtered)
            st.dataframe(view, hide_index=True, width="stretch")
            st.download_button("Baixar CSV", view.to_csv(index=False).encode("utf-8-sig"), "inventario-filtrado.csv", "text/csv", key=prefix+"csv")
            if st.button("Preparar inventário Word", key=prefix+"prepare_docx"):
                st.session_state[prefix+"docx"] = _export_docx(view)
                st.session_state[prefix+"docx_signature"] = view.to_csv(index=False)
            if st.session_state.get(prefix+"docx_signature") == view.to_csv(index=False):
                st.download_button("Baixar DOCX", st.session_state[prefix+"docx"], "inventario-filtrado.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=prefix+"download_docx")
